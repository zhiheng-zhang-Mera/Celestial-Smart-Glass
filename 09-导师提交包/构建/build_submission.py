"""Build supervisor-facing DOCX files from versioned Markdown sources.

No participant data are read. Render and visually verify before distribution.
"""
from pathlib import Path
import json
import re
import argparse
from urllib.parse import quote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PACKAGE = Path(__file__).resolve().parents[1]
REPO = PACKAGE.parent
BLACK = RGBColor(0, 0, 0)
FONT = 'Noto Sans CJK SC'
LATIN = 'Calibri'
URL_ROOT = 'https://github.com/zhiheng-zhang-Mera/Celestial-Smart-Glass/blob/main/'
KEY_PATTERN = r'(?:HCI|LI|META|EMG|APPLE|HRV|SREG|METHOD|JEV|REUSE)-\d+[A-Z]?'

def key_order(text):
    text = text.split('<!-- BEGIN GENERATED REFERENCES -->')[0]
    found=[]
    for group in re.findall(r'\[((?:'+KEY_PATTERN+r')(?:\s*,\s*'+KEY_PATTERN+r')*)\]',text):
        for key in re.findall(KEY_PATTERN,group):
            if key not in found: found.append(key)
    return found

def references():
    src=(REPO/'07-文献与证据/参考文献.md').read_text()
    out={}
    matches=list(re.finditer(r'^### ('+KEY_PATTERN+r')\s*$',src,re.M))
    for i,m in enumerate(matches):
        section=src[m.end():matches[i+1].start() if i+1<len(matches) else len(src)]
        para=next((p.strip() for p in section.split('\n\n') if p.strip() and not p.lstrip().startswith('-')), '')
        urls=re.findall(r'https?://[^\s<>）)]+',section)
        urls=[u.rstrip('。；;,') for u in urls]
        doi=next((u for u in urls if 'doi.org/' in u),None)
        url=doi or (urls[0] if urls else None)
        out[m.group(1)]=(para,url)
    return out

def materialize_refs(path):
    text=path.read_text()
    keys=key_order(text)
    entries=references()
    lines=['<!-- BEGIN GENERATED REFERENCES -->']
    for key in keys:
        if key not in entries: raise ValueError('Undefined reference '+key)
        desc,url=entries[key]
        lines.append(f'[{key}] {desc} '+(f'[来源]({url})' if url else ''))
        lines.append('')
    lines.append('<!-- END GENERATED REFERENCES -->')
    block='\n'.join(lines)
    if '<!-- REFERENCES -->' in text:
        text=text.replace('<!-- REFERENCES -->',block)
    elif '<!-- BEGIN GENERATED REFERENCES -->' in text:
        text=re.sub(r'<!-- BEGIN GENERATED REFERENCES -->.*?<!-- END GENERATED REFERENCES -->',lambda _:block,text,flags=re.S)
    else:
        raise ValueError('Missing bibliography location: '+str(path))
    path.write_text(text)
    return keys

def font(r, size=None, bold=False, mono=False):
    r.font.name='DejaVu Sans Mono' if mono else LATIN
    r.font.color.rgb=BLACK
    if size is not None:r.font.size=Pt(size)
    if bold:r.bold=True
    rf=r._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'),FONT)

def clean_heading(s):
    s=re.sub(r'[\[\]`*_#]','',s)
    s=re.sub(r'^\d+(?:\.\d+)+\s+','',s)
    return re.sub(r'[^\w\u4e00-\u9fff\s]',' ',s).strip()

def add_bookmark(p,name,number):
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(number));b.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(number))
    p._p.insert(0,b);p._p.append(e)

def hyperlink(p,label,url=None,anchor=None,bold=False,size=None):
    el=OxmlElement('w:hyperlink')
    if anchor: el.set(qn('w:anchor'),anchor)
    else: el.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),LATIN);rf.set(qn('w:hAnsi'),LATIN);rf.set(qn('w:eastAsia'),FONT);pr.append(rf)
    col=OxmlElement('w:color');col.set(qn('w:val'),'000000');pr.append(col)
    if bold: pr.append(OxmlElement('w:b'))
    if size:
        z=OxmlElement('w:sz');z.set(qn('w:val'),str(int(size*2)));pr.append(z)
    r.append(pr);t=OxmlElement('w:t');t.text='来\u2060源' if label=='来源' else label;r.append(t);el.append(r);p._p.append(el)

def resolve_link(url,source):
    if re.match(r'https?://',url):return url
    target=(source.parent/url.split('#')[0]).resolve()
    try:return URL_ROOT+quote(target.relative_to(REPO).as_posix(),safe='/')
    except ValueError:return url

def inline(p,text,source,refs,size=None):
    token=re.compile(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\[(?:'+KEY_PATTERN+r')(?:\s*,\s*'+KEY_PATTERN+r')*\])')
    for part in token.split(text):
        if not part:continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]);font(r,size,bold=True)
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]);font(r,size or 10,mono=True)
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)',part):
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',part)
            hyperlink(p,m.group(1),resolve_link(m.group(2),source),size=size)
        elif re.fullmatch(r'\[(?:'+KEY_PATTERN+r')(?:\s*,\s*'+KEY_PATTERN+r')*\]',part):
            for i,key in enumerate(re.findall(KEY_PATTERN,part)):
                if i:font(p.add_run(', '),size)
                if key in refs:hyperlink(p,f'[{refs[key]}]',anchor='ref_'+key.replace('-','_'),size=size)
                else:font(p.add_run('['+key+']'),size)
        else:
            # Plain paragraphs retain their actual wording. Only Markdown emphasis is removed.
            part=re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'\1',part)
            for j,line in enumerate(part.split('\n')):
                if j:p.add_run().add_break()
                font(p.add_run(line),size)

def table(doc,rows,source,refs):
    n=len(rows[0]);width=6.65
    if n==2: weights=[1.65,5.0]
    elif n==3:weights=[1.3,2.7,2.65]
    elif n==4 and any(x in rows[0][0] for x in ['情境','假设','场景']):weights=[3.5,1.05,1.05,1.05]
    elif n==4:weights=[1.25,1.85,1.85,1.70]
    elif n==5:weights=[1.0,1.65,1.35,1.35,1.3]
    else:weights=[width/n]*n
    if len(weights)!=n:weights=[width/n]*n
    weights=[v*width/sum(weights) for v in weights]
    t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,weights):col.width=Inches(w)
    for ri,values in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,c in enumerate(cells):
            c.width=Inches(weights[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            tcPr.append(b);m=OxmlElement('w:tcMar')
            for edge,v in [('top',90),('bottom',90),('left',100),('right',100)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');m.append(el)
            tcPr.append(m)
            p=c.paragraphs[0];p.paragraph_format.line_spacing=Pt(15);p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.space_before=Pt(2)
            val=values[ci] if ci<len(values) else ''
            if (len(val)<14 and n>=4) or ri==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,val,source,refs,size=10)
            if ri==0:
                for r in p.runs:r.bold=True
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF0F2');tcPr.append(sh)
        pr=t.rows[ri]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(3);after.paragraph_format.line_spacing=Pt(2)

def make_doc(source,output,long_form=False,toc_pages=None):
    keys=materialize_refs(source);refs={k:i+1 for i,k in enumerate(keys)}
    text=source.read_text();lines=text.splitlines()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.75)
    sec.left_margin=Inches(.925);sec.right_margin=Inches(.925)
    sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)
    normal=doc.styles['Normal'];normal.font.name=LATIN;normal.font.size=Pt(11);normal.font.color.rgb=BLACK
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    pf=normal.paragraph_format;pf.line_spacing=Pt(18);pf.space_after=Pt(6);pf.widow_control=True
    for name,size,before,after in [('Title',23,0,15),('Subtitle',14,0,12),('Heading 1',15,17,9),('Heading 2',12.5,11,6),('Heading 3',11.5,8,4)]:
        s=doc.styles[name];s.font.name=LATIN;s.font.size=Pt(size);s.font.color.rgb=BLACK;s.font.bold=True;s.font.italic=False
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        p=s.paragraph_format;p.space_before=Pt(before);p.space_after=Pt(after);p.keep_with_next=True;p.widow_control=True
        p.line_spacing=Pt(size*1.3)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        ppr=doc.styles[style_name].element.get_or_add_pPr()
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');ppr.append(snap)
    header=sec.header.paragraphs[0]
    font(header.add_run('Celestial Smart Glass   研究计划与实施方案'),9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('第 '),9)
    for instruction in ['PAGE','NUMPAGES']:
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instruction)
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);footer._p.append(field)
        font(footer.add_run(' 页  共 ' if instruction=='PAGE' else ' 页'),9)
    doc.core_properties.author='Zhiheng Zhang'
    doc.core_properties.subject='PhD research proposal and first-year project plan'
    doc.core_properties.title=lines[0].lstrip('# ').strip()
    doc.core_properties.keywords='HCI, smart glasses, accessibility, captions, user control'
    major=[clean_heading(l[3:]) for l in lines if l.startswith('## ') and re.match(r'\d+[.)]?\s',l[3:])]
    major_ids={h:'sec_'+str(i+1) for i,h in enumerate(major)}
    bookmark_id=10; toc_done=False; in_refs=False; i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line or line.startswith('<!--') or line.startswith('[返回'):
            i+=1;continue
        if line.startswith('```'):
            block=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
            for ln in block:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                font(p.add_run(ln),9.5,mono=True)
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[x.strip() for x in re.split(r'(?<!\\)\|',lines[i].strip().strip('|'))]
                if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in vals):rows.append(vals)
                i+=1
            table(doc,rows,source,refs);continue
        im=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if im:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=1.0
            picture=p.add_run().add_picture(str(source.parent/im.group(2)),width=Inches(6.4))
            picture._inline.docPr.set('descr',im.group(1))
            p.paragraph_format.keep_with_next=True
            cap=doc.add_paragraph();font(cap.add_run(im.group(1)),10);cap.paragraph_format.space_after=Pt(10)
            i+=1;continue
        hm=re.match(r'^(#{1,4})\s+(.*)',line)
        if hm:
            level=len(hm.group(1));title=clean_heading(hm.group(2))
            if level==1:p=doc.add_paragraph(style='Title')
            elif i<6 and title.startswith('博士'):p=doc.add_paragraph(style='Subtitle')
            else:
                if long_form and title in ['中文摘要','English Abstract']:
                    doc.add_page_break()
                if title=='参考文献':
                    doc.add_page_break()
                if long_form and title in major_ids and not toc_done:
                    doc.add_page_break();doc.add_paragraph('目录',style='Heading 1')
                    for h in major:
                        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
                        hyperlink(p,h,anchor=major_ids[h])
                        if toc_pages and h in toc_pages:
                            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
                            font(p.add_run('\t'+str(toc_pages[h])))
                    doc.add_page_break();toc_done=True
                p=doc.add_paragraph(style='Heading '+str(min(3,level-1)))
            font(p.add_run(title),bold=True)
            if title in major_ids:add_bookmark(p,major_ids[title],bookmark_id);bookmark_id+=1
            if title=='参考文献':in_refs=True
            elif level==2:in_refs=False
            i+=1;continue
        if in_refs:
            m=re.match(r'^\[('+KEY_PATTERN+r')\]\s+(.*)',line)
            if m:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_together=True
                p.paragraph_format.left_indent=Inches(.27);p.paragraph_format.first_line_indent=Inches(-.27)
                font(p.add_run(f'[{refs[m.group(1)]}] '),10)
                inline(p,m.group(2),source,refs,size=10)
                add_bookmark(p,'ref_'+m.group(1).replace('-','_'),bookmark_id);bookmark_id+=1
                i+=1;continue
        prefix=re.match(r'^(?:[-*]|\d+[.)、])\s+(.*)',line)
        if prefix:
            p=doc.add_paragraph(style='List Bullet' if line[0] in '-*' else 'List Number')
            inline(p,prefix.group(1),source,refs);i+=1;continue
        chunks=[lines[i]];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(?:#|\||```|!\[|<!--|\[返回|[-*]\s|\d+[.)、]\s)',lines[i].strip()):
            if in_refs and re.match(r'^\['+KEY_PATTERN,lines[i]):break
            chunks.append(lines[i]);i+=1
        joined=''
        for j,c in enumerate(chunks):
            joined+=c.rstrip()
            if j<len(chunks)-1:joined+='\n' if c.endswith('  ') else ' '
        p=doc.add_paragraph();inline(p,joined.lstrip('> '),source,refs)
    doc.save(output)
    return {'source':source.name,'output':output.name,'references':len(refs),'headings':major}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--only',choices=['main','appendix','brief']);ap.add_argument('--toc-pages',type=Path);args=ap.parse_args()
    pages=json.loads(args.toc_pages.read_text()) if args.toc_pages else {}
    jobs=[('main','完整研究计划.md','CSG_Research_Proposal_CN.docx',True),
          ('appendix','研究实施附件.md','CSG_Study_and_Implementation_Appendices_CN.docx',True),
          ('brief','研究概要.md','CSG_Research_Overview_CN.docx',False)]
    for kind,src,out,long_form in jobs:
        if args.only and args.only!=kind:continue
        print(json.dumps(make_doc(PACKAGE/src,PACKAGE/out,long_form,pages.get(kind)),ensure_ascii=False))

if __name__=='__main__':main()
